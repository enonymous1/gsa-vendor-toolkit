"""Tests for the SAM.gov integrity-records fold-in (mp-sam-data-bank).

The shipped workbook is fictional sample data with the real export's
38-column layout and two-row preamble; the assertions below pin the exact
counts the generator built into it, so a loader regression shows up here
instead of in a user's report.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from vendors_toolkit import cli
from vendors_toolkit.analysis.integrity import (
    INTEGRITY_COLUMNS,
    analyze_integrity_export,
    analyze_integrity_records,
    load_integrity_records,
    write_integrity_report,
)
from vendors_toolkit.config import IntegrityConfig

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "sample-data" / "integrity" / "sam_integrity_sample.xlsx"


def _workbook(path: Path, *, headers: list[str], rows: list[list[object]]) -> Path:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "Fictional sample"
    sheet["A2"] = "Preamble row"
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=col, value=header)
    for i, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            sheet.cell(row=4 + i, column=col, value=value)
    book.save(path)
    return path


def test_sample_workbook_present_and_fictional() -> None:
    assert SAMPLE.is_file()
    frame = load_integrity_records(SAMPLE)
    assert list(frame.columns) == list(INTEGRITY_COLUMNS)
    assert not frame["Awardee Name"].str.contains("CONFLUENCE", na=False).any()
    assert frame["Official Email"].str.endswith("@example.com").all()


def test_loader_normalizes_empties_dates_and_naics() -> None:
    frame = load_integrity_records(SAMPLE)
    assert len(frame) == 24
    # Empty cells are missing, not float NaN strings or blanks.
    assert frame["CAGE Code"].isna().sum() == 3
    assert frame["Subcontractor Name"].isna().all()
    # Dates parsed to midnight, not raw timestamp strings.
    assert str(frame["Record Date"].dtype).startswith("datetime64")
    assert frame["Record Date"].dt.hour.eq(0).all()
    # Float-formatted NAICS (541511.0 in the workbook) becomes "541511".
    assert "541511.0" not in frame["Principal NAICS Code"].dropna().tolist()
    assert set(frame["Principal NAICS Code"].dropna()) == {
        "541511",
        "541611",
        "336411",
        "541330",
    }


def test_loader_missing_file() -> None:
    with pytest.raises(FileNotFoundError, match="not found"):
        load_integrity_records(Path("/nonexistent/nope.xlsx"))


def test_loader_rejects_missing_required_column(tmp_path: Path) -> None:
    headers = [c for c in INTEGRITY_COLUMNS if c != "Agency"]
    path = _workbook(tmp_path / "bad.xlsx", headers=headers, rows=[["x"] * len(headers)])
    with pytest.raises(ValueError, match="missing required"):
        load_integrity_records(path)


def test_loader_rejects_no_header_row(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "Not an integrity export"
    path = tmp_path / "noheader.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="header row"):
        load_integrity_records(path)


def test_analysis_counts_match_sample() -> None:
    summary = analyze_integrity_records(load_integrity_records(SAMPLE))
    assert summary.total_records == 24
    assert summary.record_type_counts == {
        "Termination for Cause": 10,
        "Termination for Default": 6,
        "Administrative Agreement": 4,
        "Material Failure to Comply with Closeout Requirements": 3,
        "Non-Responsibility Determination": 1,
    }
    assert summary.most_common_agency == ("DEPT OF THE ARMY", 9)
    assert summary.most_common_naics == ("541511", 6)
    assert summary.state_counts == {"VA": 7, "MD": 6, "CA": 5, "TX": 4}
    assert summary.monthly_counts == {
        "2024-01": 3,
        "2024-03": 4,
        "2024-06": 5,
        "2024-09": 4,
        "2024-12": 4,
        "2025-02": 4,
    }
    assert summary.earliest_record_date is not None
    assert summary.latest_record_date is not None
    assert summary.earliest_record_date.year == 2024
    assert summary.latest_record_date.year == 2025
    assert "Subcontractor Name" in summary.empty_columns
    assert "Violation Against" in summary.empty_columns


def test_analysis_rejects_frame_without_required_column() -> None:
    with pytest.raises(ValueError, match="missing column"):
        analyze_integrity_records(pd.DataFrame({"Awardee Name": ["Acme"]}))


def test_report_writes_csvs_and_markdown(tmp_path: Path) -> None:
    summary = analyze_integrity_records(load_integrity_records(SAMPLE))
    written = write_integrity_report(summary, tmp_path)
    names = sorted(p.name for p in written)
    assert names == [
        "integrity_agency_counts.csv",
        "integrity_monthly_counts.csv",
        "integrity_naics_counts.csv",
        "integrity_record_type_counts.csv",
        "integrity_state_counts.csv",
        "integrity_summary.md",
    ]
    types = pd.read_csv(tmp_path / "integrity_record_type_counts.csv")
    assert list(types.columns) == ["Record Type", "count"]
    assert types.iloc[0]["Record Type"] == "Termination for Cause"
    assert int(types.iloc[0]["count"]) == 10
    markdown = (tmp_path / "integrity_summary.md").read_text(encoding="utf-8")
    assert "Total records**: 24" in markdown
    assert "DEPT OF THE ARMY (9)" in markdown


def test_integrity_config_defaults() -> None:
    cfg = IntegrityConfig()
    assert cfg.input_dirname == "integrity"
    assert cfg.output_dirname == "integrity"
    assert cfg.header_search_rows == 5


def test_engine_discovers_input_and_writes_reports(cfg, tmp_path: Path) -> None:
    input_dir = tmp_path / "integrity"
    input_dir.mkdir()
    (input_dir / "export.xlsx").write_bytes(SAMPLE.read_bytes())
    written = analyze_integrity_export(cfg)
    assert any(p.name == "integrity_summary.md" for p in written)
    assert (tmp_path / "output" / "integrity" / "integrity_summary.md").is_file()


def test_engine_missing_input_raises(cfg) -> None:
    with pytest.raises(FileNotFoundError, match="No .xlsx export"):
        analyze_integrity_export(cfg)


def test_cli_integrity_runs_once_with_all_variants(monkeypatch, tmp_path: Path) -> None:
    """The integrity export is variant-independent: one run, not two."""
    input_dir = tmp_path / "integrity"
    input_dir.mkdir()
    (input_dir / "export.xlsx").write_bytes(SAMPLE.read_bytes())
    rc = cli.main(["--root", str(tmp_path), "integrity"])
    assert rc == 0
    assert (tmp_path / "output" / "integrity" / "integrity_summary.md").is_file()


def test_cli_integrity_missing_input_returns_1(tmp_path: Path) -> None:
    rc = cli.main(["--root", str(tmp_path), "integrity"])
    assert rc == 1


def test_cli_integrity_explicit_input_and_output(tmp_path: Path, capsys) -> None:
    out_dir = tmp_path / "reports"
    rc = cli.main(
        [
            "--root",
            str(tmp_path),
            "integrity",
            "--input",
            str(SAMPLE),
            "--output",
            str(out_dir),
        ]
    )
    assert rc == 0
    assert (out_dir / "integrity_summary.md").is_file()
    assert "integrity report ->" in capsys.readouterr().out


def test_parser_has_integrity_stage() -> None:
    args = cli.build_parser().parse_args(["integrity", "--input", "x.xlsx"])
    assert args.stage == "integrity"
    assert args.input == "x.xlsx"
